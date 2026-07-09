from pathlib import Path
import textwrap

import openpyxl
from docx import Document
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "public" / "assets" / "ai"
SOURCE = Path(r"D:\ai创意策划 曾璇 19980811942\测评反馈")


def font(size, bold=False):
    candidates = [
        r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
    ]
    for item in candidates:
        path = Path(item)
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


FONTS = {
    "hero": font(54, True),
    "title": font(36, True),
    "sub": font(23, True),
    "body": font(20),
    "small": font(17),
    "tiny": font(15),
}


def wrap_text(draw, text, font_obj, max_width):
    lines = []
    for paragraph in str(text).splitlines():
        paragraph = paragraph.strip()
        if not paragraph:
            lines.append("")
            continue
        current = ""
        for char in paragraph:
            candidate = current + char
            if draw.textbbox((0, 0), candidate, font=font_obj)[2] <= max_width:
                current = candidate
            else:
                if current:
                    lines.append(current)
                current = char
        if current:
            lines.append(current)
    return lines


def draw_wrapped(draw, xy, text, font_obj, fill, max_width, line_gap=8):
    x, y = xy
    for line in wrap_text(draw, text, font_obj, max_width):
        draw.text((x, y), line, font=font_obj, fill=fill)
        y += font_obj.size + line_gap
    return y


def panel(draw, box, outline=(55, 255, 183, 86), fill=(9, 18, 17)):
    draw.rounded_rectangle(box, radius=22, fill=fill, outline=outline, width=1)


def load_script_rows():
    wb = openpyxl.load_workbook(SOURCE / "脚本.xlsx", data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        cells = list(row)
        if cells and cells[0] is None:
            cells = cells[1:]
        time, shot, desc, caption, *_ = cells + [None] * 4
        if not time or not shot or not desc:
            continue
        rows.append(
            {
                "time": str(time).strip(),
                "shot": str(shot).strip(),
                "desc": " ".join(str(desc).split()),
                "caption": " ".join(str(caption or "").split()),
            }
        )
    return rows


def load_brief():
    doc = Document(SOURCE / "AI创意策划 曾璇 19980811942.docx")
    return [p.text.strip() for p in doc.paragraphs if p.text.strip()]


def background(size):
    img = Image.new("RGB", size, "#050706")
    draw = ImageDraw.Draw(img, "RGBA")
    w, h = size
    for y in range(h):
        t = y / h
        r = int(5 + 5 * t)
        g = int(7 + 24 * t)
        b = int(8 + 13 * t)
        draw.line((0, y, w, y), fill=(r, g, b, 255))
    draw.ellipse((-260, -280, 520, 520), fill=(42, 255, 181, 26))
    draw.ellipse((w - 460, h - 460, w + 260, h + 260), fill=(96, 255, 203, 18))
    for x in range(80, w, 160):
        draw.line((x, 0, x, h), fill=(255, 255, 255, 10), width=1)
    return img, draw


def save_overview(rows, brief):
    img, draw = background((1800, 1180))
    draw.text((90, 84), "AI 视频脚本与分镜策划", font=FONTS["hero"], fill="#f3fff8")
    draw.text((92, 154), "Whiteout Survival / 冰雪题材 AI 广告测试", font=FONTS["body"], fill="#8fffd1")

    panel(draw, (90, 230, 610, 470))
    draw.text((128, 266), "任务要求", font=FONTS["title"], fill="#ffffff")
    brief_text = "围绕冰雪题材展开，前半段用 AI 呈现强视觉冲击与心流体验，后半段衔接游戏玩法，可加入广告语与解说文案。"
    draw_wrapped(draw, (128, 330), brief_text, FONTS["body"], "#c7d8d2", 415)

    stat_cards = [("80s", "成片脚本时长"), (str(len(rows)), "镜头段落"), ("4", "脚本字段"), ("AI+游戏", "混剪结构")]
    x = 660
    for value, label in stat_cards:
        panel(draw, (x, 230, x + 250, 470), fill=(8, 20, 17))
        draw.text((x + 32, 284), value, font=FONTS["hero"], fill="#3cffb5")
        draw.text((x + 34, 372), label, font=FONTS["small"], fill="#c7d8d2")
        x += 280

    draw.text((90, 560), "叙事拆解", font=FONTS["title"], fill="#ffffff")
    milestones = [
        ("0-15s", "末日冰封", "地球、城市、避难所与两类玩家的死亡结局。"),
        ("15-27s", "重生回档", "回到末日前 24 小时，用强钩子制造继续观看。"),
        ("27-40s", "策略改写", "零氪抢生路，氪佬清内鬼，人物动机对应玩法。"),
        ("40-61s", "极寒追击", "用寒潮、车队、熔炉中心把视觉冲击推到高潮。"),
        ("61-80s", "玩法衔接", "升级熔炉、资源生产、联盟集结、下载转化。"),
    ]
    y = 630
    for idx, (time, title, body) in enumerate(milestones, start=1):
        panel(draw, (90, y, 1710, y + 86), fill=(7, 15, 14))
        draw.text((125, y + 24), f"{idx:02d}", font=FONTS["sub"], fill="#3cffb5")
        draw.text((205, y + 20), time, font=FONTS["sub"], fill="#ffffff")
        draw.text((360, y + 20), title, font=FONTS["sub"], fill="#ffffff")
        draw_wrapped(draw, (570, y + 21), body, FONTS["small"], "#b7c9c2", 1040)
        y += 106

    img.save(OUT / "ai-script-overview.jpg", quality=92)


def save_scenes(rows):
    img, draw = background((1800, 1500))
    draw.text((90, 78), "关键镜头脚本样张", font=FONTS["hero"], fill="#f3fff8")
    draw.text((92, 148), "从时间、镜头语言、画面描述到字幕/旁白的完整拆分", font=FONTS["body"], fill="#8fffd1")

    anchors = [0, 2, 6, 8, 10, 13, 16, len(rows) - 1]
    selected = []
    for index in anchors:
        safe_index = max(0, min(index, len(rows) - 1))
        if safe_index not in selected:
            selected.append(safe_index)
    while len(selected) < 8 and len(selected) < len(rows):
        candidate = len(selected)
        if candidate not in selected:
            selected.append(candidate)
    card_w, card_h = 790, 270
    positions = [(90, 235), (920, 235), (90, 535), (920, 535), (90, 835), (920, 835), (90, 1135), (920, 1135)]
    for pos, row_index in zip(positions, selected):
        row = rows[row_index]
        x, y = pos
        panel(draw, (x, y, x + card_w, y + card_h), fill=(8, 17, 16))
        draw.text((x + 30, y + 26), row["time"], font=FONTS["title"], fill="#3cffb5")
        draw_wrapped(draw, (x + 190, y + 30), row["shot"], FONTS["sub"], "#ffffff", 540, 7)
        desc = textwrap.shorten(row["desc"], width=96, placeholder="...")
        caption = textwrap.shorten(row["caption"], width=58, placeholder="...")
        draw_wrapped(draw, (x + 30, y + 100), desc, FONTS["small"], "#c6d5cf", 700, 7)
        draw.line((x + 30, y + 202, x + card_w - 30, y + 202), fill=(60, 255, 181, 55), width=1)
        draw_wrapped(draw, (x + 30, y + 220), caption, FONTS["tiny"], "#9dffda", 700, 5)

    img.save(OUT / "ai-script-scenes.jpg", quality=92)


def save_workflow(rows):
    img, draw = background((1800, 1080))
    draw.text((90, 78), "AI 内容生产链路", font=FONTS["hero"], fill="#f3fff8")
    draw.text((92, 148), "把创意脚本转成可生成、可剪辑、可衔接玩法的短视频执行文件", font=FONTS["body"], fill="#8fffd1")

    steps = [
        ("01", "Brief 拆解", "冰雪题材、视觉冲击、AI 画面、玩法衔接、广告语与旁白。"),
        ("02", "剧情钩子", "末日冰封、双角色死亡、同时重生，制造短视频前 3 秒吸引力。"),
        ("03", "镜头语言", "极远景、航拍、推镜、低机位、匹配转场、游戏 UI 推入。"),
        ("04", "画面提示", "每个镜头写出场景、人物、道具、运动、光线和情绪，用作生成基础。"),
        ("05", "字幕旁白", "同步补充字幕、旁白和人物台词，保证剪辑时信息闭环。"),
        ("06", "成片衔接", "AI 电影化画面过渡到熔炉升级、资源生产、联盟集结和下载按钮。"),
    ]
    x_positions = [90, 650, 1210]
    y_positions = [250, 610]
    for idx, (num, title, body) in enumerate(steps):
        x = x_positions[idx % 3]
        y = y_positions[idx // 3]
        panel(draw, (x, y, x + 500, y + 260), fill=(7, 17, 15))
        draw.text((x + 34, y + 32), num, font=FONTS["hero"], fill="#3cffb5")
        draw.text((x + 145, y + 50), title, font=FONTS["title"], fill="#ffffff")
        draw_wrapped(draw, (x + 34, y + 134), body, FONTS["body"], "#c7d8d2", 420, 10)

    img.save(OUT / "ai-video-workflow.jpg", quality=92)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    rows = load_script_rows()
    brief = load_brief()
    save_overview(rows, brief)
    save_scenes(rows)
    save_workflow(rows)
    print(f"generated {len(rows)} script rows into {OUT}")


if __name__ == "__main__":
    main()
